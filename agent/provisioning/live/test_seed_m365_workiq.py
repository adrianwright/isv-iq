"""Targeted tests for the Microsoft 365 Work IQ proof-of-concept-content seeder.

Coverage:
  * payload generation (deterministic subjects, required evidence tokens, transactionId, time zone),
  * idempotent selection and upsert decisions (skip/send, create/patch),
  * Graph retry and error handling (429/5xx backoff, permission errors, give-up),
  * no token output (client repr, result summary, device-code print).

These are offline unit tests: every Graph interaction runs through httpx.MockTransport, and the MSAL
app is a stub. No network, no real sign-in.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import date
from pathlib import Path

import httpx
import pytest

import seed_m365_workiq as seed
from seed_m365_workiq import (
    EMAIL_SUBJECT,
    CRCL_EVENT_SUBJECT,
    PI_EVENT_SUBJECT,
    PI_TODO_TASK_TITLE,
    TODO_LIST_NAME,
    TODO_TASK_TITLE,
    DeviceCodeAuthError,
    GraphClient,
    GraphError,
    GraphPermissionError,
    SeedContext,
    acquire_token_device_code,
    build_calendar_events,
    build_email_message,
    build_todo_task,
    build_todo_tasks,
    run_seed,
    seed_email,
    seed_events,
    seed_todo,
    select_existing_event,
    select_existing_list,
    select_existing_message,
    select_existing_task,
)

TOKEN = "super-secret-access-token-value"


def _ctx() -> SeedContext:
    return SeedContext(user_address="analyst@contoso.onmicrosoft.com")


def _client(handler, **kwargs) -> GraphClient:
    transport = httpx.MockTransport(handler)
    http_client = httpx.Client(transport=transport, base_url="https://graph.microsoft.com/v1.0")
    return GraphClient(TOKEN, http_client=http_client, sleep=lambda _s: None, **kwargs)


# --------------------------------------------------------------------------------------------------
# Payload generation
# --------------------------------------------------------------------------------------------------
def test_email_subject_is_deterministic_and_versioned() -> None:
    assert EMAIL_SUBJECT == seed.EMAIL_SUBJECT
    assert seed.SUBJECT_VERSION in EMAIL_SUBJECT
    assert "TASK-1042-CRCL" in EMAIL_SUBJECT


def test_build_email_message_addresses_self_and_carries_all_evidence() -> None:
    ctx = _ctx()
    message = build_email_message(ctx)

    assert message["message"]["subject"] == EMAIL_SUBJECT
    assert message["message"]["toRecipients"][0]["emailAddress"]["address"] == ctx.user_address

    body = message["message"]["body"]["content"]
    for token in [
        "TASK-1042-CRCL",
        "Dana Whitfield",
        "COORD-01",
        "TB-2026-06-30",
        "48 mL/min",
        "50 mL/min",
        "Dr. Priya Anand",
        "prior-platinum",
        "RQ-THORACIC",
        "Awaiting screening",
        "2026-07-08",
        "ECOG 1",
        "CKD-EPI",
        "several days old",
        "repeat required",
        "Amendment 2",
        "AMD-2",
        "line 1 / first-line",
        "Partial response",
        "Synthetic",
        "Not clinical decision support",
    ]:
        assert token in body, f"expected '{token}' in email body"


def test_build_calendar_events_are_deterministic_with_transaction_ids() -> None:
    ctx = _ctx()
    events = build_calendar_events(ctx)
    assert [event["subject"] for event in events] == [CRCL_EVENT_SUBJECT, PI_EVENT_SUBJECT]

    crcl, pi = events
    assert crcl["start"]["timeZone"] == "Eastern Standard Time"
    assert crcl["start"]["dateTime"].startswith(ctx.crcl_due_date)
    assert pi["start"]["dateTime"].startswith(ctx.pi_slot_date)
    assert crcl["transactionId"] and pi["transactionId"]
    assert crcl["transactionId"] != pi["transactionId"]
    # Stable across builds so a rerun before delivery cannot double-create.
    assert crcl["transactionId"] == build_calendar_events(ctx)[0]["transactionId"]
    assert "TB-2026-06-30" in crcl["body"]["content"]
    assert "Dr. Priya Anand" in pi["body"]["content"]
    assert all(
        criterion_id in pi["body"]["content"]
        for criterion_id in [
            "NCT99004324-DX",
            "NCT99004324-BIO",
            "NCT99004324-PS",
            "NCT99004324-REN",
            "NCT99004324-RX",
        ]
    )


def test_build_todo_task_is_open_and_identifies_owner_and_ids() -> None:
    task = build_todo_task(_ctx())
    assert task["title"] == TODO_TASK_TITLE
    assert task["status"] == "notStarted"
    assert task["dueDateTime"]["dateTime"].startswith("2026-07-08")
    assert task["dueDateTime"]["timeZone"] == "Eastern Standard Time"

    body = task["body"]["content"]
    for token in ["Dana Whitfield", "PT-1042", "NCT99004324", "TB-2026-06-30", "Dr. Priya Anand", "Not clinical decision support"]:
        assert token in body, f"expected '{token}' in todo body"


def test_build_todo_tasks_materializes_crcl_and_pi_review_as_discrete_tasks() -> None:
    tasks = build_todo_tasks(_ctx())
    assert [task["title"] for task in tasks] == [TODO_TASK_TITLE, PI_TODO_TASK_TITLE]
    assert [task["dueDateTime"]["dateTime"][:10] for task in tasks] == [
        "2026-07-08",
        "2026-07-09",
    ]
    assert "Dana Whitfield (COORD-01)" in tasks[0]["body"]["content"]
    assert "Dr. Priya Anand (PI-01)" in tasks[1]["body"]["content"]
    assert "TASK-1042-PI-REVIEW" in tasks[1]["body"]["content"]


def test_delegated_scopes_cover_collaboration_and_copilot_without_application_permissions() -> None:
    assert set(seed.DELEGATED_SCOPES) == {
        "User.Read",
        "Mail.ReadBasic",
        "Mail.Send",
        "Calendars.ReadWrite",
        "Tasks.ReadWrite",
        "Files.ReadWrite.All",
        "Files.Read.All",
        "Sites.Read.All",
        "Chat.Read",
        "ChatMessage.Send",
    }


def test_powershell_provisioner_scope_allowlist_matches_python_seeder() -> None:
    wrapper = Path(seed.__file__).with_name("provision_m365_workiq.ps1").read_text(encoding="utf-8")
    match = re.search(r"\$ScopeValues\s*=\s*@\((.*?)\)", wrapper, flags=re.DOTALL)
    assert match, "PowerShell scope allowlist was not found"
    wrapper_scopes = set(re.findall(r'"([^"]+)"', match.group(1)))
    assert wrapper_scopes == set(seed.DELEGATED_SCOPES)
    assert f'$WorkIqScopeValue = "{seed.WORK_IQ_SCOPE.rsplit("/", 1)[-1]}"' in wrapper


def test_build_docx_transcripts_are_valid_word_packages_with_semantic_evidence() -> None:
    documents = seed.build_transcript_documents(_ctx())
    assert [document.filename for document in documents] == [
        seed.TUMOR_BOARD_TRANSCRIPT_FILENAME,
        seed.SCREENING_HUDDLE_TRANSCRIPT_FILENAME,
    ]

    for document in documents:
        package = seed.build_docx(document)
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            assert set(archive.namelist()) == {
                "[Content_Types].xml",
                "_rels/.rels",
                "word/document.xml",
            }
            xml = archive.read("word/document.xml").decode("utf-8")
        for token in ["PT-1042", "NCT99004324", "Synthetic", "Not clinical decision support"]:
            assert token in xml


def test_transcript_comparison_ignores_sharepoint_package_metadata() -> None:
    expected = seed.build_docx(seed.build_transcript_documents(_ctx())[0])
    modified = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(expected)) as source, zipfile.ZipFile(
        modified, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for name in source.namelist():
            target.writestr(name, source.read(name))
        target.writestr("docProps/sharepoint.xml", b"<sharepoint>metadata</sharepoint>")

    assert modified.getvalue() != expected
    assert seed.transcript_content_matches(modified.getvalue(), expected)
    assert not seed.transcript_content_matches(b"not-a-docx", expected)


def test_build_productivity_files_contains_word_excel_and_powerpoint_evidence() -> None:
    from openpyxl import load_workbook

    data_dir = Path(seed.__file__).resolve().parents[3] / "data" / "work"
    files = seed.build_productivity_files(_ctx(), data_dir)
    assert [item.filename for item in files] == [
        seed.PRODUCTIVITY_MEMO_FILENAME,
        seed.PRODUCTIVITY_TRACKER_FILENAME,
        seed.PRODUCTIVITY_DECK_FILENAME,
    ]
    assert [item.retrieval_mode for item in files] == ["semantic", "lexical", "semantic"]

    extracted = {
        item.filename: {
            ".docx": seed.docx_text,
            ".xlsx": seed.xlsx_text,
            ".pptx": seed.pptx_text,
        }[Path(item.filename).suffix](item.content)
        for item in files
    }
    for filename, text in extracted.items():
        assert "PT-1042" in text, filename
        assert "NCT99004324" in text, filename
        assert "Synthetic" in text, filename
    assert "Screening Tracker" in extracted[seed.PRODUCTIVITY_TRACKER_FILENAME]
    assert "Thoracic Trials Weekly Operations" in extracted[seed.PRODUCTIVITY_DECK_FILENAME]
    for filename in [seed.PRODUCTIVITY_MEMO_FILENAME, seed.PRODUCTIVITY_DECK_FILENAME]:
        for criterion_id in [
            "NCT99004324-DX",
            "NCT99004324-BIO",
            "NCT99004324-PS",
            "NCT99004324-REN",
            "NCT99004324-RX",
        ]:
            assert criterion_id in extracted[filename]

    tracker_file = next(
        item for item in files if item.filename == seed.PRODUCTIVITY_TRACKER_FILENAME
    )
    workbook = load_workbook(io.BytesIO(tracker_file.content), read_only=True, data_only=True)
    try:
        tracker = workbook["Screening Tracker"]
        headers = [cell.value for cell in tracker[1]]
        pt1042 = next(row for row in tracker.iter_rows(values_only=True) if row[0] == "PT-1042")
        row = dict(zip(headers, pt1042, strict=True))
        assert row["Owner"] == "Dana Whitfield"
        assert row["Due Date"] == "2026-07-08"
        assert row["Open Tasks"] == "TASK-1042-CRCL, TASK-1042-PI-REVIEW"
    finally:
        workbook.close()


def test_productivity_comparison_ignores_office_package_metadata() -> None:
    data_dir = Path(seed.__file__).resolve().parents[3] / "data" / "work"
    for artifact in seed.build_productivity_files(_ctx(), data_dir):
        modified = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(artifact.content)) as source, zipfile.ZipFile(
            modified, "w", compression=zipfile.ZIP_DEFLATED
        ) as target:
            for name in source.namelist():
                target.writestr(name, source.read(name))
            target.writestr("docProps/sharepoint-seed.xml", b"<sharepoint>metadata</sharepoint>")
        assert seed.office_content_matches(
            artifact.filename, modified.getvalue(), artifact.content
        )
        assert not seed.office_content_matches(
            artifact.filename, b"not-an-office-file", artifact.content
        )


def test_build_self_chat_messages_are_marked_and_non_impersonating() -> None:
    messages = seed.build_self_chat_messages(_ctx())
    assert len(messages) == len(seed.SELF_CHAT_MARKERS) == 3
    for marker, message in zip(seed.SELF_CHAT_MARKERS, messages, strict=True):
        content = message["body"]["content"]
        assert marker in content
        assert "Synthetic" in content
        assert message["body"]["contentType"] == "text"


def test_build_seed_context_reads_data_work_artifacts() -> None:
    data_dir = Path(seed.__file__).resolve().parents[3] / "data" / "work"
    ctx = seed.build_seed_context(
        data_dir, "analyst@contoso.onmicrosoft.com", reference_date=date(2026, 7, 1)
    )
    assert ctx.owner_name == "Dana Whitfield"
    assert ctx.owner_id == "COORD-01"
    assert ctx.crcl_due_date == "2026-07-08"
    assert ctx.referral_status == "Awaiting screening"
    assert ctx.pi_slot_date == "2026-07-07"
    assert ctx.pi_task_id == "TASK-1042-PI-REVIEW"
    assert ctx.pi_owner_id == "PI-01"
    assert (ctx.amendment_label, ctx.amendment_id) == ("Amendment 2", "AMD-2")
    assert "prior first-line platinum may be permitted with PI confirmation" in ctx.amendment_summary
    assert (ctx.prior_therapy_line, ctx.prior_therapy_response) == ("1", "Partial response")
    assert {criterion_id.rsplit("-", 1)[-1] for criterion_id, _ in ctx.protocol_criteria} == {
        "DX",
        "BIO",
        "PS",
        "REN",
        "RX",
    }


def test_build_seed_context_shifts_stale_action_dates_as_one_ordered_block() -> None:
    from openpyxl import load_workbook

    data_dir = Path(seed.__file__).resolve().parents[3] / "data" / "work"
    ctx = seed.build_seed_context(
        data_dir, "analyst@contoso.onmicrosoft.com", reference_date=date(2026, 7, 22)
    )
    assert (ctx.pi_slot_date, ctx.crcl_due_date, ctx.pi_review_due_date) == (
        "2026-07-23",
        "2026-07-24",
        "2026-07-25",
    )
    assert ctx.crcl_last_date == "2026-06-18"

    tracker_file = next(
        item
        for item in seed.build_productivity_files(ctx, data_dir)
        if item.filename == seed.PRODUCTIVITY_TRACKER_FILENAME
    )
    workbook = load_workbook(io.BytesIO(tracker_file.content), read_only=True, data_only=True)
    try:
        task_rows = {
            row[0]: row
            for row in workbook["Open Tasks"].iter_rows(min_row=2, values_only=True)
        }
        assert task_rows["TASK-1042-CRCL"][6] == "2026-07-24"
        assert task_rows["TASK-1042-PI-REVIEW"][6] == "2026-07-25"
        patient_review = next(
            row
            for row in workbook["PI Availability"].iter_rows(min_row=2, values_only=True)
            if row[5] == "Patient Reviews"
        )
        assert patient_review[2:5] == ("2026-07-23", "09:00", "10:00")
    finally:
        workbook.close()


def test_existing_calendar_event_pins_seed_version_schedule() -> None:
    data_dir = Path(seed.__file__).resolve().parents[3] / "data" / "work"
    unshifted = seed.build_seed_context(
        data_dir,
        "analyst@contoso.onmicrosoft.com",
        reference_date=date.min,
    )

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path.endswith("/me/events")
        return httpx.Response(
            200,
            request=request,
            json={
                "value": [
                    {
                        "id": "event-1",
                        "subject": seed.CRCL_EVENT_SUBJECT,
                        "start": {
                            "dateTime": "2026-07-24T09:00:00.0000000",
                            "timeZone": seed.EVENT_TIME_ZONE,
                        },
                    }
                ]
            },
        )

    reference_date = seed.resolve_seed_reference_date(
        _client(handler),
        unshifted,
        fallback_date=date(2026, 7, 24),
    )
    assert reference_date == date(2026, 7, 22)

    pinned = seed.build_seed_context(
        data_dir,
        "analyst@contoso.onmicrosoft.com",
        reference_date=reference_date,
    )
    assert (pinned.pi_slot_date, pinned.crcl_due_date, pinned.pi_review_due_date) == (
        "2026-07-23",
        "2026-07-24",
        "2026-07-25",
    )


def test_missing_calendar_event_uses_current_seed_reference() -> None:
    reference_date = seed.resolve_seed_reference_date(
        _client(lambda request: httpx.Response(200, request=request, json={"value": []})),
        _ctx(),
        fallback_date=date(2026, 7, 24),
    )
    assert reference_date == date(2026, 7, 24)


def test_workiq_answer_requires_all_scenario_critical_evidence() -> None:
    ctx = _ctx()
    answer = (
        "For PT-1042 and NCT99004324, Dana Whitfield (COORD-01) owns TASK-1042-CRCL, "
        "due 2026-07-08. The tumor board TB-2026-06-30 recorded CrCl 48 mL/min using "
        "CKD-EPI against the 50 mL/min "
        "threshold and said likely eligible pending repeat CrCl. Dr. Priya Anand (PI-01) "
        "owns TASK-1042-PI-REVIEW, due 2026-07-09, for the prior-platinum clarification. "
        "Keep RQ-THORACIC "
        "in Awaiting screening and use the 2026-07-07 Patient Reviews slot."
    )
    checks = seed.validate_workiq_answer(answer, ctx)
    assert checks and all(checks.values())


def test_workiq_answer_reports_missing_evidence() -> None:
    with pytest.raises(seed.WorkIQVerificationError, match="pi_review_owner"):
        seed.validate_workiq_answer(
            "PT-1042 NCT99004324 TASK-1042-CRCL Dana Whitfield 2026-07-08 48 mL/min "
            "50 mL/min CKD-EPI TB-2026-06-30 likely eligible TASK-1042-PI-REVIEW "
            "2026-07-09 prior platinum "
            "RQ-THORACIC Awaiting screening 2026-07-07",
            _ctx(),
        )


def test_build_seed_context_reports_missing_required_fabric_sources(tmp_path: Path) -> None:
    work_dir = tmp_path / "data" / "work"
    work_dir.mkdir(parents=True)
    with pytest.raises(seed.SeedError, match="amendments.csv"):
        seed.build_seed_context(
            work_dir,
            "analyst@contoso.onmicrosoft.com",
            reference_date=date(2026, 7, 1),
        )


def test_boilerplate_work_iq_questions_have_maintainable_seed_coverage() -> None:
    ctx = _ctx()
    data_dir = Path(seed.__file__).resolve().parents[3] / "data" / "work"
    events = build_calendar_events(ctx)
    todo_tasks = build_todo_tasks(ctx)
    productivity = {
        item.filename: {
            ".docx": seed.docx_text,
            ".xlsx": seed.xlsx_text,
            ".pptx": seed.pptx_text,
        }[Path(item.filename).suffix](item.content)
        for item in seed.build_productivity_files(ctx, data_dir)
    }
    artifacts = {
        "mail": build_email_message(ctx)["message"]["body"]["content"],
        "pi_calendar": events[1]["body"]["content"],
        "todo_all": "\n".join(task["body"]["content"] for task in todo_tasks),
        "pi_todo": todo_tasks[1]["body"]["content"],
        "chat": "\n".join(
            message["body"]["content"] for message in seed.build_self_chat_messages(ctx)
        ),
        "memo": productivity[seed.PRODUCTIVITY_MEMO_FILENAME],
        "deck": productivity[seed.PRODUCTIVITY_DECK_FILENAME],
    }
    coverage = {
        "Trial Eligibility": {
            "artifacts": ("mail", "memo"),
            "tokens": ("likely eligible", "repeat CrCl", "PI confirmation", "NCT99004324"),
        },
        "Screening Check": {
            "artifacts": ("mail", "todo_all"),
            "tokens": ("48 mL/min", "repeat required", "Awaiting screening"),
        },
        "Evidence": {
            "artifacts": ("memo", "pi_calendar", "deck"),
            "tokens": tuple(f"NCT99004324-{suffix}" for suffix in ("DX", "BIO", "PS", "REN", "RX")),
            "each_artifact": True,
        },
        "Workflow": {
            "artifacts": ("todo_all",),
            "tokens": (
                "TASK-1042-CRCL",
                "TASK-1042-PI-REVIEW",
                "Dana Whitfield",
                "Dr. Priya Anand",
                "2026-07-08",
                "2026-07-09",
            ),
        },
        "Protocol": {
            "artifacts": ("mail", "pi_calendar", "pi_todo", "memo", "deck"),
            "tokens": (
                "Amendment 2",
                "AMD-2",
                "line-of-therapy",
                "prior first-line platinum may be permitted with PI confirmation",
                "line 1 / first-line",
                "Partial response",
            ),
            "each_artifact": True,
        },
        "Data Gaps": {
            "artifacts": ("mail", "todo_all", "chat", "memo", "deck"),
            "tokens": (
                "ECOG 1",
                "CKD-EPI",
                "48 mL/min",
                "2026-06-18",
                ">= 50 mL/min",
                "several days old",
                "repeat required",
            ),
            "each_artifact": True,
        },
    }

    question_library = (
        Path(seed.__file__).resolve().parents[3] / "apps" / "web" / "src" / "questionLibrary.ts"
    ).read_text(encoding="utf-8")
    work_iq_categories = set(re.findall(r"category:\s*'([^']+)'", question_library)) - {
        "External Context"
    }
    assert set(coverage) == work_iq_categories

    for category, requirement in coverage.items():
        texts = [artifacts[name] for name in requirement["artifacts"]]
        haystacks = texts if requirement.get("each_artifact") else ["\n".join(texts)]
        for artifact_text in haystacks:
            for token in requirement["tokens"]:
                assert token in artifact_text, (
                    f"{category} seed coverage missing {token!r} in "
                    f"{requirement['artifacts']}"
                )




# --------------------------------------------------------------------------------------------------
# Idempotent selection helpers
# --------------------------------------------------------------------------------------------------
def test_selection_helpers_match_case_insensitively_and_miss_cleanly() -> None:
    assert select_existing_message([{"subject": EMAIL_SUBJECT.upper()}], EMAIL_SUBJECT) is not None
    assert select_existing_message([{"subject": "unrelated"}], EMAIL_SUBJECT) is None
    assert select_existing_event([{"subject": CRCL_EVENT_SUBJECT}], CRCL_EVENT_SUBJECT) is not None
    assert select_existing_list([{"displayName": " AMC IQ Demo "}], TODO_LIST_NAME) is not None
    assert select_existing_task([{"title": TODO_TASK_TITLE}], TODO_TASK_TITLE) is not None
    assert select_existing_task([{"title": PI_TODO_TASK_TITLE}], PI_TODO_TASK_TITLE) is not None
    assert select_existing_task([], TODO_TASK_TITLE) is None


# --------------------------------------------------------------------------------------------------
# Idempotent upsert decisions (against MockTransport)
# --------------------------------------------------------------------------------------------------
def test_seed_email_skips_when_message_already_present() -> None:
    calls: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(f"{request.method} {request.url.path}")
        if request.url.path.endswith("/me/messages"):
            return httpx.Response(200, json={"value": [{"id": "m1", "subject": EMAIL_SUBJECT}]})
        raise AssertionError(f"unexpected call {request.method} {request.url.path}")

    result = seed_email(_client(handler), _ctx())
    assert result["action"] == "skipped"
    assert not any("sendMail" in call for call in calls)


def test_seed_email_sends_when_absent() -> None:
    calls: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(f"{request.method} {request.url.path}")
        if request.url.path.endswith("/me/messages"):
            return httpx.Response(200, json={"value": []})
        if request.url.path.endswith("/me/sendMail"):
            payload = json.loads(request.content)
            assert payload["message"]["subject"] == EMAIL_SUBJECT
            return httpx.Response(202)
        raise AssertionError(f"unexpected call {request.method} {request.url.path}")

    result = seed_email(_client(handler), _ctx())
    assert result["action"] == "sent"
    assert any("sendMail" in call for call in calls)


def test_seed_events_creates_then_patches() -> None:
    def make_handler(existing: bool):
        def handler(request: httpx.Request) -> httpx.Response:
            path = request.url.path
            if request.method == "GET" and path.endswith("/me/events"):
                value = [{"id": "e-existing", "subject": _subject_from_filter(request)}] if existing else []
                return httpx.Response(200, json={"value": value})
            if request.method == "POST" and path.endswith("/me/events"):
                return httpx.Response(201, json={"id": "e-new"})
            if request.method == "PATCH" and "/me/events/" in path:
                return httpx.Response(200, json={"id": path.rsplit("/", 1)[-1]})
            raise AssertionError(f"unexpected call {request.method} {path}")

        return handler

    created = seed_events(_client(make_handler(existing=False)), _ctx())
    assert [entry["action"] for entry in created] == ["created", "created"]

    patched = seed_events(_client(make_handler(existing=True)), _ctx())
    assert [entry["action"] for entry in patched] == ["patched", "patched"]


def test_seed_events_patch_omits_transaction_id() -> None:
    seen_bodies: list[dict] = []

    def handler(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        if request.method == "GET" and path.endswith("/me/events"):
            return httpx.Response(200, json={"value": [{"id": "e1", "subject": _subject_from_filter(request)}]})
        if request.method == "PATCH" and "/me/events/" in path:
            seen_bodies.append(json.loads(request.content))
            return httpx.Response(200, json={"id": "e1"})
        raise AssertionError(f"unexpected call {request.method} {path}")

    seed_events(_client(handler), _ctx())
    assert seen_bodies, "expected a PATCH body"
    assert all("transactionId" not in body for body in seen_bodies)


def test_seed_todo_creates_list_and_task_when_absent() -> None:
    created_titles: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        path, method = request.url.path, request.method
        if method == "GET" and path.endswith("/me/todo/lists"):
            return httpx.Response(200, json={"value": []})
        if method == "POST" and path.endswith("/me/todo/lists"):
            assert json.loads(request.content)["displayName"] == TODO_LIST_NAME
            return httpx.Response(201, json={"id": "list-1", "displayName": TODO_LIST_NAME})
        if method == "GET" and path.endswith("/tasks"):
            return httpx.Response(200, json={"value": []})
        if method == "POST" and path.endswith("/tasks"):
            payload = json.loads(request.content)
            assert payload["status"] == "notStarted"
            created_titles.append(payload["title"])
            return httpx.Response(201, json={"id": f"task-{len(created_titles)}"})
        raise AssertionError(f"unexpected call {method} {path}")

    result = seed_todo(_client(handler), _ctx())
    assert result["list_action"] == "created"
    assert [task["action"] for task in result["tasks"]] == ["created", "created"]
    assert created_titles == [TODO_TASK_TITLE, PI_TODO_TASK_TITLE]


def test_seed_todo_reuses_list_and_patches_existing_task() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        path, method = request.url.path, request.method
        if method == "GET" and path.endswith("/me/todo/lists"):
            return httpx.Response(200, json={"value": [{"id": "list-1", "displayName": TODO_LIST_NAME}]})
        if method == "GET" and path.endswith("/tasks"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {"id": "task-1", "title": TODO_TASK_TITLE},
                        {"id": "task-2", "title": PI_TODO_TASK_TITLE},
                    ]
                },
            )
        if method == "PATCH" and "/tasks/" in path:
            return httpx.Response(200, json={"id": path.rsplit("/", 1)[-1]})
        raise AssertionError(f"unexpected call {method} {path}")

    result = seed_todo(_client(handler), _ctx())
    assert result["list_action"] == "reused"
    assert [task["action"] for task in result["tasks"]] == ["patched", "patched"]


def test_seed_sharepoint_transcripts_creates_folders_and_uploads_raw_docx() -> None:
    folder_ids = iter(["folder-amc", "folder-transcripts"])
    uploaded: list[bytes] = []

    def handler(request: httpx.Request) -> httpx.Response:
        path, method = request.url.path, request.method
        if method == "GET" and path.endswith("/sites/root"):
            return httpx.Response(
                200,
                json={"id": "site-1", "displayName": "Communication site", "webUrl": "https://example.sharepoint.com"},
            )
        if method == "GET" and path.endswith("/sites/site-1/drive"):
            return httpx.Response(200, json={"id": "drive-1", "name": "Documents", "webUrl": "https://example.sharepoint.com/Documents"})
        if method == "GET" and path.endswith("/drives/drive-1/root"):
            return httpx.Response(200, json={"id": "root-1"})
        if method == "GET" and "/drives/drive-1/root:/" in path:
            return httpx.Response(404, json={"error": {"code": "itemNotFound"}})
        if method == "POST" and "/children" in path:
            payload = json.loads(request.content)
            folder_id = next(folder_ids)
            return httpx.Response(
                201,
                json={"id": folder_id, "name": payload["name"], "folder": {}, "webUrl": f"https://example.sharepoint.com/{folder_id}"},
            )
        if method == "PUT" and path.endswith(":/content"):
            assert request.headers["Content-Type"].startswith(
                "application/vnd.openxmlformats-officedocument"
            )
            assert request.headers["If-None-Match"] == "*"
            uploaded.append(request.content)
            return httpx.Response(
                201,
                json={"id": f"file-{len(uploaded)}", "webUrl": f"https://example.sharepoint.com/file-{len(uploaded)}.docx"},
            )
        raise AssertionError(f"unexpected call {method} {path}")

    result = seed.seed_sharepoint_transcripts(_client(handler), _ctx())
    assert result["folder_path"] == seed.SHAREPOINT_FOLDER_PATH
    assert [item["action"] for item in result["files"]] == ["created", "created"]
    assert len(uploaded) == 2
    assert all(zipfile.is_zipfile(io.BytesIO(content)) for content in uploaded)


def test_seed_sharepoint_transcript_refuses_to_overwrite_changed_file(monkeypatch) -> None:
    monkeypatch.setattr(
        seed,
        "ensure_sharepoint_transcript_folder",
        lambda _graph: {
            "site_name": "Communication site",
            "site_web_url": "https://example.sharepoint.com",
            "drive_id": "drive-1",
            "drive_name": "Documents",
            "folder_path": seed.SHAREPOINT_FOLDER_PATH,
            "folder_web_url": "https://example.sharepoint.com/folder",
        },
    )

    def handler(request: httpx.Request) -> httpx.Response:
        path, method = request.url.path, request.method
        if method == "GET" and "/root:/" in path:
            return httpx.Response(200, json={"id": "existing", "name": "existing.docx", "webUrl": "https://example/existing"})
        if method == "GET" and path.endswith("/items/existing/content"):
            return httpx.Response(200, content=b"user-edited-content")
        raise AssertionError(f"unexpected call {method} {path}")

    with pytest.raises(GraphError, match="Refusing to overwrite"):
        seed.seed_sharepoint_transcripts(_client(handler), _ctx())


def test_seed_sharepoint_transcript_recovers_after_precondition_response(monkeypatch) -> None:
    document = seed.build_transcript_documents(_ctx())[0]
    expected_content = seed.build_docx(document)
    monkeypatch.setattr(seed, "build_transcript_documents", lambda _ctx: (document,))
    monkeypatch.setattr(
        seed,
        "ensure_sharepoint_transcript_folder",
        lambda _graph: {
            "site_id": "site-1",
            "site_name": "Communication site",
            "site_web_url": "https://example.sharepoint.com",
            "drive_id": "drive-1",
            "drive_name": "Documents",
            "folder_path": seed.SHAREPOINT_FOLDER_PATH,
            "folder_web_url": "https://example.sharepoint.com/folder",
        },
    )
    path_reads = {"count": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        path, method = request.url.path, request.method
        if method == "GET" and "/root:/" in path:
            path_reads["count"] += 1
            if path_reads["count"] == 1:
                return httpx.Response(404, json={"error": {"code": "itemNotFound"}})
            return httpx.Response(
                200,
                json={"id": "recovered", "name": document.filename, "webUrl": "https://example/recovered.docx"},
            )
        if method == "PUT" and path.endswith(":/content"):
            return httpx.Response(412, json={"error": {"code": "preconditionFailed"}})
        if method == "GET" and path.endswith("/items/recovered/content"):
            return httpx.Response(200, content=expected_content)
        raise AssertionError(f"unexpected call {method} {path}")

    result = seed.seed_sharepoint_transcripts(_client(handler), _ctx())
    assert result["files"][0]["action"] == "recovered"


def test_seed_teams_self_chat_skips_existing_and_sends_missing() -> None:
    payloads = seed.build_self_chat_messages(_ctx())
    sent: list[dict] = []

    def handler(request: httpx.Request) -> httpx.Response:
        path, method = request.url.path, request.method
        if method == "GET" and path.endswith("/chats/48:notes/messages"):
            return httpx.Response(
                200,
                json={"value": [{"id": "existing", "body": {"content": payloads[0]["body"]["content"]}}]},
            )
        if method == "POST" and path.endswith("/chats/48:notes/messages"):
            sent.append(json.loads(request.content))
            return httpx.Response(201, json={"id": f"sent-{len(sent)}", "body": sent[-1]["body"]})
        raise AssertionError(f"unexpected call {method} {path}")

    result = seed.seed_teams_self_chat(_client(handler), _ctx())
    assert [item["action"] for item in result["messages"]] == ["skipped", "sent", "sent"]
    assert len(sent) == 2


def test_seed_teams_self_chat_retries_429_without_duplicate_probe() -> None:
    post_calls = {"count": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        path, method = request.url.path, request.method
        if method == "GET" and path.endswith("/chats/48:notes/messages"):
            return httpx.Response(200, json={"value": []})
        if method == "POST" and path.endswith("/chats/48:notes/messages"):
            post_calls["count"] += 1
            if post_calls["count"] == 1:
                return httpx.Response(
                    429,
                    headers={"Retry-After": "0"},
                    json={"error": {"code": "TooManyRequests"}},
                )
            payload = json.loads(request.content)
            return httpx.Response(
                201, json={"id": f"sent-{post_calls['count']}", "body": payload["body"]}
            )
        raise AssertionError(f"unexpected call {method} {path}")

    result = seed.seed_teams_self_chat(_client(handler), _ctx())
    assert [item["action"] for item in result["messages"]] == ["sent", "sent", "sent"]
    assert post_calls["count"] == 4


def test_list_teams_self_chat_messages_follows_next_link() -> None:
    calls = {"count": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        calls["count"] += 1
        if calls["count"] == 1:
            return httpx.Response(
                200,
                json={
                    "value": [{"id": "m1", "body": {"content": "first"}}],
                    "@odata.nextLink": "https://graph.microsoft.com/v1.0/chats/48:notes/messages?page=2",
                },
            )
        return httpx.Response(200, json={"value": [{"id": "m2", "body": {"content": "second"}}]})

    messages = seed.list_teams_self_chat_messages(_client(handler))
    assert [message["id"] for message in messages] == ["m1", "m2"]


def test_teams_self_chat_not_found_is_explicit() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(404, json={"error": {"code": "NotFound"}})

    with pytest.raises(GraphError, match="Open your own name in Teams"):
        seed.list_teams_self_chat_messages(_client(handler))


def test_copilot_retrieval_reports_indexed_only_when_both_transcripts_match() -> None:
    filenames = [
        seed.TUMOR_BOARD_TRANSCRIPT_FILENAME,
        seed.SCREENING_HUDDLE_TRANSCRIPT_FILENAME,
    ]

    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        assert payload["dataSource"] == "sharePoint"
        assert "Filename:" in payload["filterExpression"]
        return httpx.Response(
            200,
            json={
                "retrievalHits": [
                    {
                        "webUrl": f"https://example.sharepoint.com/{filename}",
                        "extracts": [{"text": f"PT-1042 and NCT99004324 evidence from {filename}"}],
                    }
                    for filename in filenames
                ]
            },
        )

    indexed = seed.verify_copilot_sharepoint_retrieval(
        _client(handler),
        {
            "folder_web_url": "https://example.sharepoint.com/folder",
            "files": [
                {"filename": filename, "web_url": f"https://example.sharepoint.com/{filename}"}
                for filename in filenames
            ],
        },
    )
    assert indexed["status"] == "indexed"
    assert indexed["evidence_found"] is True


def test_copilot_retrieval_reports_pending_when_index_has_no_hits() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"retrievalHits": []})

    pending = seed.verify_copilot_sharepoint_retrieval(
        _client(handler),
        {
            "folder_web_url": "https://example.sharepoint.com/folder",
            "files": [
                {
                    "filename": seed.TUMOR_BOARD_TRANSCRIPT_FILENAME,
                    "web_url": f"https://example.sharepoint.com/{seed.TUMOR_BOARD_TRANSCRIPT_FILENAME}",
                },
                {
                    "filename": seed.SCREENING_HUDDLE_TRANSCRIPT_FILENAME,
                    "web_url": f"https://example.sharepoint.com/{seed.SCREENING_HUDDLE_TRANSCRIPT_FILENAME}",
                },
            ],
        },
    )
    assert pending["status"] == "pending"
    assert pending["matched_files"] == []


def test_copilot_retrieval_rejects_same_named_files_from_another_location() -> None:
    filenames = [
        seed.TUMOR_BOARD_TRANSCRIPT_FILENAME,
        seed.SCREENING_HUDDLE_TRANSCRIPT_FILENAME,
    ]

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "retrievalHits": [
                    {
                        "webUrl": f"https://example.sharepoint.com/other/{filename}",
                        "extracts": [{"text": "PT-1042 NCT99004324"}],
                    }
                    for filename in filenames
                ]
            },
        )

    result = seed.verify_copilot_sharepoint_retrieval(
        _client(handler),
        {
            "folder_web_url": "https://example.sharepoint.com/folder",
            "files": [
                {"filename": filename, "web_url": f"https://example.sharepoint.com/folder/{filename}"}
                for filename in filenames
            ],
        },
    )
    assert result["status"] == "pending"
    assert result["matched_files"] == []


def test_copilot_productivity_retrieval_requires_semantic_word_and_powerpoint() -> None:
    filenames = [
        seed.PRODUCTIVITY_MEMO_FILENAME,
        seed.PRODUCTIVITY_TRACKER_FILENAME,
        seed.PRODUCTIVITY_DECK_FILENAME,
    ]

    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        assert "Path:" in payload["filterExpression"]
        return httpx.Response(
            200,
            json={
                "retrievalHits": [
                    {
                        "webUrl": f"https://example.sharepoint.com/productivity/{filename}",
                        "extracts": [{"text": f"PT-1042 NCT99004324 evidence in {filename}"}],
                    }
                    for filename in filenames
                ]
            },
        )

    result = seed.verify_copilot_productivity_retrieval(
        _client(handler),
        {
            "folder_web_url": "https://example.sharepoint.com/productivity",
            "files": [
                {
                    "filename": filename,
                    "web_url": f"https://example.sharepoint.com/productivity/{filename}",
                }
                for filename in filenames
            ],
        },
    )
    assert result["status"] == "indexed"
    assert set(result["semantic_expected_files"]) == {
        seed.PRODUCTIVITY_MEMO_FILENAME,
        seed.PRODUCTIVITY_DECK_FILENAME,
    }
    assert result["lexical_expected_files"] == [seed.PRODUCTIVITY_TRACKER_FILENAME]


def test_copilot_productivity_retrieval_does_not_require_excel_for_semantic_status() -> None:
    semantic_files = [seed.PRODUCTIVITY_MEMO_FILENAME, seed.PRODUCTIVITY_DECK_FILENAME]
    all_files = semantic_files + [seed.PRODUCTIVITY_TRACKER_FILENAME]

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "retrievalHits": [
                    {
                        "webUrl": f"https://example.sharepoint.com/productivity/{filename}",
                        "extracts": [{"text": "PT-1042 NCT99004324 operations plan"}],
                    }
                    for filename in semantic_files
                ]
            },
        )

    result = seed.verify_copilot_productivity_retrieval(
        _client(handler),
        {
            "folder_web_url": "https://example.sharepoint.com/productivity",
            "files": [
                {
                    "filename": filename,
                    "web_url": f"https://example.sharepoint.com/productivity/{filename}",
                }
                for filename in all_files
            ],
        },
    )
    assert result["status"] == "indexed"
    assert seed.PRODUCTIVITY_TRACKER_FILENAME not in result["matched_files"]


# --------------------------------------------------------------------------------------------------
# Graph retry / error handling
# --------------------------------------------------------------------------------------------------
def test_retries_on_429_then_succeeds_and_honors_retry_after() -> None:
    attempts = {"n": 0}
    slept: list[float] = []

    def handler(request: httpx.Request) -> httpx.Response:
        attempts["n"] += 1
        if attempts["n"] == 1:
            return httpx.Response(429, headers={"Retry-After": "7"}, json={"error": {"code": "tooManyRequests"}})
        return httpx.Response(200, json={"ok": True})

    transport = httpx.MockTransport(handler)
    http_client = httpx.Client(transport=transport, base_url="https://graph.microsoft.com/v1.0")
    client = GraphClient(TOKEN, http_client=http_client, sleep=slept.append)

    response = client.request("GET", "/me")
    assert response.json() == {"ok": True}
    assert attempts["n"] == 2
    assert slept == [7.0]


def test_retries_on_503_then_succeeds() -> None:
    attempts = {"n": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        attempts["n"] += 1
        if attempts["n"] < 3:
            return httpx.Response(503, json={"error": {"code": "serviceUnavailable"}})
        return httpx.Response(200, json={"ok": True})

    assert _client(handler).request("GET", "/me").json() == {"ok": True}
    assert attempts["n"] == 3


def test_gives_up_after_max_attempts() -> None:
    attempts = {"n": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        attempts["n"] += 1
        return httpx.Response(500, json={"error": {"code": "internalServerError"}})

    with pytest.raises(GraphError, match="after 3 attempts"):
        _client(handler, max_attempts=3).request("GET", "/me")
    assert attempts["n"] == 3


def test_permission_error_is_explicit_and_not_retried() -> None:
    attempts = {"n": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        attempts["n"] += 1
        return httpx.Response(403, json={"error": {"code": "accessDenied", "message": "Insufficient privileges"}})

    with pytest.raises(GraphPermissionError, match="permission or unavailable-workload"):
        _client(handler).request("GET", "/me")
    assert attempts["n"] == 1


def test_client_error_is_raised_without_retry() -> None:
    attempts = {"n": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        attempts["n"] += 1
        return httpx.Response(400, json={"error": {"code": "badRequest", "message": "bad filter"}})

    with pytest.raises(GraphError, match="HTTP 400"):
        _client(handler).request("GET", "/me")
    assert attempts["n"] == 1


def test_retries_on_network_error_then_succeeds() -> None:
    attempts = {"n": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        attempts["n"] += 1
        if attempts["n"] == 1:
            raise httpx.ConnectError("boom", request=request)
        return httpx.Response(200, json={"ok": True})

    assert _client(handler).request("GET", "/me").json() == {"ok": True}
    assert attempts["n"] == 2


def test_network_error_is_wrapped_as_graph_error_after_max_attempts() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectTimeout("timeout", request=request)

    with pytest.raises(GraphError, match="after 2 attempts"):
        _client(handler, max_attempts=2).request("GET", "/me")


def test_requests_carry_the_bearer_token() -> None:
    seen: dict[str, str | None] = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["auth"] = request.headers.get("Authorization")
        return httpx.Response(200, json={"ok": True})

    _client(handler).request("GET", "/me")
    assert seen["auth"] == f"Bearer {TOKEN}"


# --------------------------------------------------------------------------------------------------
# No token output
# --------------------------------------------------------------------------------------------------
def test_graph_client_repr_never_leaks_token() -> None:
    client = GraphClient(TOKEN, http_client=httpx.Client(base_url="https://graph.microsoft.com/v1.0"))
    assert TOKEN not in repr(client)


def test_run_seed_summary_contains_no_token(monkeypatch) -> None:
    monkeypatch.setattr(seed, "seed_email", lambda _graph, _ctx: {"action": "sent"})
    monkeypatch.setattr(seed, "seed_events", lambda _graph, _ctx: [{"action": "created"}])
    monkeypatch.setattr(seed, "seed_todo", lambda _graph, _ctx: {"action": "created"})
    monkeypatch.setattr(
        seed,
        "seed_sharepoint_transcripts",
        lambda _graph, _ctx: {"files": [{"action": "created"}], "folder_web_url": "https://example/folder"},
    )
    monkeypatch.setattr(
        seed,
        "seed_sharepoint_productivity",
        lambda _graph, _ctx, _data_dir: {
            "files": [{"action": "created"}],
            "folder_web_url": "https://example/productivity",
        },
    )
    monkeypatch.setattr(
        seed,
        "seed_teams_self_chat",
        lambda _graph, _ctx: {"messages": [{"action": "sent"}]},
    )
    monkeypatch.setattr(
        seed,
        "verify_copilot_sharepoint_retrieval",
        lambda _graph, _sharepoint: {"status": "pending"},
    )
    monkeypatch.setattr(
        seed,
        "verify_copilot_productivity_retrieval",
        lambda _graph, _sharepoint: {"status": "pending"},
    )

    result = run_seed(_client(lambda request: httpx.Response(500)), _ctx())
    assert result["status"] == "ok"
    assert result["mail"]["action"] == "sent"
    assert result["sharepoint"]["files"][0]["action"] == "created"
    assert result["sharepoint_productivity"]["files"][0]["action"] == "created"
    assert result["copilot_retrieval"]["status"] == "pending"
    assert result["copilot_productivity_retrieval"]["status"] == "pending"
    assert TOKEN not in json.dumps(result)


def test_run_seed_reports_copilot_license_block_without_losing_seed_result(monkeypatch) -> None:
    monkeypatch.setattr(seed, "seed_email", lambda _graph, _ctx: {"action": "skipped"})
    monkeypatch.setattr(seed, "seed_events", lambda _graph, _ctx: [{"action": "patched"}])
    monkeypatch.setattr(seed, "seed_todo", lambda _graph, _ctx: {"action": "patched"})
    monkeypatch.setattr(
        seed,
        "seed_sharepoint_transcripts",
        lambda _graph, _ctx: {"files": [{"action": "skipped"}], "folder_web_url": "https://example/folder"},
    )
    monkeypatch.setattr(
        seed,
        "seed_sharepoint_productivity",
        lambda _graph, _ctx, _data_dir: {
            "files": [{"action": "skipped"}],
            "folder_web_url": "https://example/productivity",
        },
    )
    monkeypatch.setattr(
        seed,
        "seed_teams_self_chat",
        lambda _graph, _ctx: {"messages": [{"action": "skipped"}]},
    )

    def blocked(_graph, _sharepoint):
        raise GraphPermissionError("User does not have valid license")

    monkeypatch.setattr(seed, "verify_copilot_sharepoint_retrieval", blocked)
    result = run_seed(_client(lambda request: httpx.Response(500)), _ctx())
    assert result["status"] == "ok"
    assert result["copilot_retrieval"]["status"] == "blocked"
    assert result["copilot_productivity_retrieval"]["status"] == "blocked"
    assert "valid license" in result["copilot_retrieval"]["error"]


def test_device_code_uses_silent_cached_token_before_interactive_flow() -> None:
    class _CachedApp:
        def get_accounts(self):
            return [{"home_account_id": "account-1"}]

        def acquire_token_silent(self, scopes, account):
            assert list(scopes) == list(seed.DELEGATED_SCOPES)
            assert account["home_account_id"] == "account-1"
            return {"access_token": TOKEN}

        def initiate_device_flow(self, scopes):
            raise AssertionError("device flow should not start when the cache can refresh")

    assert acquire_token_device_code("tenant", "client", app=_CachedApp()) == TOKEN


def test_device_code_prints_message_only_and_returns_token() -> None:
    printed: list[str] = []

    class _App:
        def initiate_device_flow(self, scopes):
            assert list(scopes) == list(seed.DELEGATED_SCOPES)
            return {"user_code": "AB12CD34", "message": "Go to https://microsoft.com/devicelogin and enter AB12CD34"}

        def acquire_token_by_device_flow(self, flow):
            return {"access_token": TOKEN}

    token = acquire_token_device_code("tenant", "client", app=_App(), print_fn=printed.append)
    assert token == TOKEN
    assert printed == ["Go to https://microsoft.com/devicelogin and enter AB12CD34"]
    assert all(TOKEN not in line for line in printed)


def test_device_code_failure_is_explicit() -> None:
    class _BadApp:
        def initiate_device_flow(self, scopes):
            return {"error": "invalid_client", "error_description": "app not found"}

    with pytest.raises(DeviceCodeAuthError, match="app not found"):
        acquire_token_device_code("tenant", "client", app=_BadApp())


def test_device_code_no_token_is_explicit() -> None:
    class _NoTokenApp:
        def initiate_device_flow(self, scopes):
            return {"user_code": "X", "message": "sign in"}

        def acquire_token_by_device_flow(self, flow):
            return {"error": "authorization_declined", "error_description": "user declined"}

    with pytest.raises(DeviceCodeAuthError, match="user declined"):
        acquire_token_device_code("tenant", "client", app=_NoTokenApp())


def _subject_from_filter(request: httpx.Request) -> str:
    """Extract the subject value from a `$filter=subject eq '...'` query for echoing back."""
    filter_value = dict(request.url.params).get("$filter", "")
    start = filter_value.find("'")
    end = filter_value.rfind("'")
    return filter_value[start + 1 : end].replace("''", "'") if start != -1 and end > start else ""
